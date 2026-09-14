"""Unit tests for the export pipeline: markdown → DOCX / PDF, verified by
reading the produced files back with python-docx and pypdf."""

from __future__ import annotations

import io

from app.agent.demo_fixture import DEMO_PROJECT_TITLE
from app.agent.nodes.finalize import compose_final_markdown
from app.services.exporter import markdown_to_docx_bytes, markdown_to_pdf_bytes

SECTIONS = [
    {
        "title": "Project Summary",
        "order_index": 1,
        "content_md": "## Project Summary\n\nWe will deploy **40 sensors** and train 60 wardens.\n\n- Open hardware\n- Multilingual alerts",
    },
    {
        "title": "Statement of Need",
        "order_index": 2,
        "content_md": "## Statement of Need\n\nFlood risk is rated in the 92nd percentile [1].",
    },
    {
        "title": "Evaluation Plan",
        "order_index": 3,
        "content_md": "## Evaluation Plan\n\n1. Baseline\n2. Midline\n3. Endline",
    },
]
FINDINGS = [{"title": "Risk Atlas", "url": "https://mock.tavily.local/atlas"}]
ABSTRACT = "A twenty-four month community resilience project targeting measurable gains in warning lead time."
TITLE = "Community Flood-Sensor Network and Resident Response Program"


def demo_markdown() -> str:
    return compose_final_markdown(TITLE, ABSTRACT, SECTIONS, FINDINGS)


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------
def test_docx_contains_title_headings_and_lists():
    from docx import Document

    data = markdown_to_docx_bytes(demo_markdown(), title=TITLE, organization="Riverbend")
    document = Document(io.BytesIO(data))
    text = "\n".join(p.text for p in document.paragraphs)
    assert TITLE in text
    assert "Project Summary" in text and "Statement of Need" in text
    assert "Open hardware" in text  # bullet list items survive
    assert "40 sensors" in text
    assert "[1] Risk Atlas — https://mock.tavily.local/atlas" in text
    headings = [p.text for p in document.paragraphs if p.style.name.startswith("Heading")]
    assert any("Abstract" in h for h in headings)
    assert any("References" in h for h in headings)


def test_docx_bold_runs_rendered():
    from docx import Document

    data = markdown_to_docx_bytes("**Bold claim** and plain.", title="T")
    document = Document(io.BytesIO(data))
    bold_texts = [r.text for p in document.paragraphs for r in p.runs if r.bold]
    assert "Bold claim" in bold_texts


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------
def test_pdf_is_valid_and_contains_content():
    import pypdf

    data = markdown_to_pdf_bytes(demo_markdown(), title=TITLE, organization="Riverbend")
    assert data[:5] == b"%PDF-"
    reader = pypdf.PdfReader(io.BytesIO(data))
    # cover page + body pages
    assert len(reader.pages) >= 2
    cover = " ".join((reader.pages[0].extract_text() or "").split())
    assert "Community Flood-Sensor Network and Resident Response Program" in cover
    assert "Archimedes" in cover  # attribution line
    body = "\n".join((p.extract_text() or "") for p in reader.pages[1:])
    assert "Statement of Need" in body
    assert "References" in body
    assert "mock.tavily.local/atlas" in body


def test_pdf_escapes_markup_characters():
    data = markdown_to_pdf_bytes("## Heading\n\n5 < 6 & 7 > 2", title="T & Co")
    import pypdf

    text = "\n".join((p.extract_text() or "") for p in pypdf.PdfReader(io.BytesIO(data)).pages)
    assert "5 < 6 & 7 > 2" in text  # escaped, not swallowed


# ---------------------------------------------------------------------------
# Assembled markdown (the export router's input)
# ---------------------------------------------------------------------------
def test_composed_markdown_includes_all_parts():
    md = demo_markdown()
    assert md.startswith(f"# {TITLE}")
    assert ABSTRACT.split()[0] in md
    assert "## Table of Contents" in md
    assert DEMO_PROJECT_TITLE  # fixture sanity
